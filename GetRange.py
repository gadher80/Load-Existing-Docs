"""
Code : Extracting a range of data from raw excel sheet and create multiple docx suign different row inputs
Authour : Hardik Gadher
Date : 30/05/2022
Location: Berlin

"""

import pandas as pd
from docx import Document
import os
import tkinter as tk
from tkinter import simpledialog

#initializing a dataframe from raw file
df01 = pd.read_excel('Raw data.xlsm')
print(f'There are total {len(df01)} rows\n')

#Input window
ROOT = tk.Tk()
ROOT.withdraw()
Start = simpledialog.askstring(title="Test",prompt='Start ?')
Start = int(Start)
Stop = simpledialog.askstring(title="Test",prompt='Stop?')
Stop = int(Stop)

#iterrate over row values for a given range
for index in range(Start - 1,Stop) :
    Probe = Document('Router.docx')
    table = Probe.tables[0]
    df01['Router No'] = df01['Router No'].astype('int')

    #assigining a value to a cell of docx from a cell of raw excel (df)
    table.cell(0,1).text = str(df01['Router No'][index].item())
    table.cell(1,1).text = str(df01['Order Number'][index].item())
    table.cell(0,3).text = str(df01['Description'][index]) + ' / ' + str(df01['Size'][index])
    table.cell(0,5).text = str(df01['Article No'][index])
    table.cell(1,5).text = str(df01['Product Id'][index]) + ', ' + str(df01['Article No'][index]) + ', ' + str(df01['Revision'][index])
    table.cell(0,7).text = str((df01['Lot No'][index]).item()) + ', ' + str((df01['Quantity'][index]).item())
    table.cell(1,7).text = str(df01['Expiry Date'][index])
    Router_number = table.cell(0,1).text

    #creating a folder and save file if not exsist
    if not os.path.exists(f'C:\\Personal\\Aufgaben\\Code\\1\\{Router_number}') :
        path = os.path.join('C:\\Personal\\Aufgaben\\Code\\1',f'{Router_number}')
        os.mkdir(path)
    Probe.save(f'C:\\Personal\\Aufgaben\\Code\\1\\{Router_number}\\{Router_number}.docx')

